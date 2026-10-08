"""Выгрузка в Excel + архив с чеками."""
import io
import re
import zipfile
from collections import OrderedDict
from collections.abc import Awaitable, Callable
from dataclasses import dataclass

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .calc import EXPENSE_KINDS, Rates, ShiftCalc
from .db import STATUS_LABELS, Database
from .views import compute, hm, ru_date, shift_km

# file_id → (содержимое, расширение с точкой)
FileFetcher = Callable[[str, str], Awaitable[tuple[bytes, str]]]

MONEY_FMT = '#,##0.00 "₽"'
HEAD_FILL = PatternFill("solid", fgColor="DDE7F3")
TOTAL_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

SHIFT_COLUMNS = [
    ("Дата", 11, None),
    ("Водитель", 22, None),
    ("Авто", 18, None),
    ("Начало", 8, None),
    ("Конец", 8, None),
    ("Часов", 7, "0.00"),
    ("Перераб., ч", 10, "0.00"),
    ("Пробег, км", 10, "0"),
    ("Сверх, км", 9, "0"),
    ("Смена", 12, MONEY_FMT),
    ("Переработка", 12, MONEY_FMT),
    ("Км сверх", 11, MONEY_FMT),
    ("Платные дороги", 13, MONEY_FMT),
    ("Парковки", 11, MONEY_FMT),
    ("Топливо", 11, MONEY_FMT),
    ("Прочее", 10, MONEY_FMT),
    ("ИТОГО", 13, MONEY_FMT),
    ("Чеков", 7, "0"),
    ("Статус", 12, None),
]
SUMMARY_COLUMNS = [
    ("Контрагент", 28, None),
    ("Водителей", 10, "0"),
    ("Смен", 7, "0"),
    ("Перераб., ч", 10, "0.00"),
    ("Сверх, км", 9, "0"),
    ("Смены", 13, MONEY_FMT),
    ("Переработка", 12, MONEY_FMT),
    ("Км сверх", 11, MONEY_FMT),
    ("Платные дороги", 13, MONEY_FMT),
    ("Парковки", 11, MONEY_FMT),
    ("Топливо", 11, MONEY_FMT),
    ("Прочее", 10, MONEY_FMT),
    ("ИТОГО", 14, MONEY_FMT),
]


@dataclass
class ExportResult:
    filename: str
    data: bytes
    shifts: int
    total: float
    missing_receipts: int


def _safe(text: str, limit: int = 40) -> str:
    text = re.sub(r'[\\/:*?"<>|\[\]\n\r\t]+', " ", text or "").strip()
    return re.sub(r"\s+", "_", text)[:limit] or "без_имени"


def _sheet_title(name: str, used: set[str]) -> str:
    base = re.sub(r"[\\/:*?\[\]]", " ", name or "Без контрагента").strip()[:28] or "Лист"
    title, i = base, 2
    while title.lower() in used:
        title = f"{base[:25]} ({i})"
        i += 1
    used.add(title.lower())
    return title


def _header(ws, row: int, columns) -> None:
    for col, (title, width, _) in enumerate(columns, start=1):
        c = ws.cell(row=row, column=col, value=title)
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
        c.border = BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _row(ws, row: int, values, columns, total: bool = False) -> None:
    for col, value in enumerate(values, start=1):
        c = ws.cell(row=row, column=col, value=value)
        fmt = columns[col - 1][2]
        if fmt and isinstance(value, (int, float)):
            c.number_format = fmt
        c.border = BORDER
        if total:
            c.font = Font(bold=True)
            c.fill = TOTAL_FILL


def _money_parts(calc: ShiftCalc) -> list[float]:
    e = calc.expenses
    return [calc.shift_pay, calc.overtime_pay, calc.km_pay,
            e["toll"], e["parking"], e["fuel"], e["other"], calc.total]


def period_label(date_from: str | None, date_to: str | None) -> str:
    if not date_from and not date_to:
        return "весь период"
    a = ru_date(date_from) if date_from else "…"
    b = ru_date(date_to) if date_to else "…"
    return f"{a} – {b}"


async def build_export(
    db: Database,
    project_id: int,
    date_from: str | None = None,
    date_to: str | None = None,
    contractor: str | None = None,
    fetch_file: FileFetcher | None = None,
    statuses: tuple[str, ...] = ("submitted", "approved"),
) -> ExportResult:
    project = await db.get_project(project_id)
    shifts = await db.shifts_for_export(project_id, date_from, date_to, statuses, contractor)
    expenses = await db.expenses_for_shifts([s["id"] for s in shifts])
    members: dict[int, dict] = {}

    # Группировка по контрагенту: одному ИП/фирме — один лист и одна строка сводки
    groups: "OrderedDict[str, list[tuple[dict, ShiftCalc]]]" = OrderedDict()
    for s in shifts:
        if s["member_id"] not in members:
            members[s["member_id"]] = await db.get_member(s["member_id"])
        calc, _ = await compute(db, s, members[s["member_id"]], project, expenses[s["id"]])
        groups.setdefault(s["contractor"] or "Без контрагента", []).append((s, calc))

    period = period_label(date_from, date_to)
    wb = Workbook()
    summary = wb.active
    summary.title = "Сводка"
    summary["A1"] = f"Проект: {project['name']}"
    summary["A1"].font = Font(bold=True, size=14)
    summary["A2"] = f"Период: {period}"
    rates = Rates.merged(project)
    summary["A3"] = (
        f"Условия проекта: смена {rates.shift_rate:g} ₽ / {rates.shift_hours:g} ч, "
        f"переработка {rates.overtime_rate:g} ₽/ч, "
        f"км: {rates.km_included:g} включено, сверх — {rates.km_rate:g} ₽/км"
    )
    summary["A4"] = "Статусы: " + ", ".join(STATUS_LABELS[s] for s in statuses)
    _header(summary, 6, SUMMARY_COLUMNS)

    receipts_rows: list[list] = []
    zip_files: list[tuple[str, bytes]] = []
    missing = 0
    used_titles = {"сводка", "чеки"}
    row_s = 7
    grand = [0.0] * 11
    for contractor_name, items in groups.items():
        ws = wb.create_sheet(_sheet_title(contractor_name, used_titles))
        ws["A1"] = f"{contractor_name} · {project['name']} · {period}"
        ws["A1"].font = Font(bold=True, size=12)
        _header(ws, 3, SHIFT_COLUMNS)
        r = 4
        sums = [0.0] * 12
        receipts_count = 0
        for s, calc in items:
            exps = expenses[s["id"]]
            km = shift_km(s)
            values = [
                ru_date(s["work_date"]), s["driver_name"], s["car"],
                hm(s["start_at"]), hm(s["end_at"]),
                round(calc.hours, 2), calc.overtime_hours, km or 0, calc.km_over,
                *_money_parts(calc),
                sum(1 for e in exps if e["file_id"]),
                STATUS_LABELS.get(s["status"], s["status"]),
            ]
            _row(ws, r, values, SHIFT_COLUMNS)
            for i, v in enumerate(values[5:17]):
                sums[i] += v
            receipts_count += values[17]
            r += 1

            for e in exps:
                link = None
                if e["file_id"]:
                    if fetch_file:
                        content, ext = await fetch_file(e["file_id"], e["file_kind"] or "photo")
                        link = (
                            f"Чеки/{_safe(contractor_name)}/"
                            f"{s['work_date']}_{_safe(s['driver_name'], 25)}_"
                            f"{_safe(EXPENSE_KINDS.get(e['kind'], e['kind']), 15)}_"
                            f"{e['amount']:g}_{e['id']}{ext}"
                        )
                        zip_files.append((link, content))
                else:
                    missing += 1
                receipts_rows.append([
                    ru_date(s["work_date"]), contractor_name, s["driver_name"], s["car"],
                    EXPENSE_KINDS.get(e["kind"], e["kind"]), e["amount"],
                    link or ("нет чека" if not e["file_id"] else "есть в боте"),
                ])

        drivers = len({s["member_id"] for s, _ in items})
        _row(ws, r, ["ИТОГО", f"{drivers} вод.", "", "", "", *sums,
                     receipts_count, ""], SHIFT_COLUMNS, total=True)
        ws.auto_filter.ref = f"A3:{get_column_letter(len(SHIFT_COLUMNS))}{r - 1}"

        money_cols = sums[4:12]
        _row(summary, row_s, [contractor_name, drivers, len(items), sums[1], sums[3],
                              *money_cols], SUMMARY_COLUMNS)
        row_s += 1
        grand[0] += len(items)
        grand[1] += sums[1]
        grand[2] += sums[3]
        for i, v in enumerate(money_cols):
            grand[3 + i] += v

    total_drivers = len({s["member_id"] for s in shifts})
    _row(summary, row_s, ["ИТОГО", total_drivers, int(grand[0]), *grand[1:]],
         SUMMARY_COLUMNS, total=True)

    receipts_cols = [("Дата", 11, None), ("Контрагент", 26, None), ("Водитель", 22, None),
                     ("Авто", 18, None), ("Тип", 16, None), ("Сумма", 12, MONEY_FMT),
                     ("Чек", 60, None)]
    wr = wb.create_sheet("Чеки")
    _header(wr, 1, receipts_cols)
    for i, values in enumerate(receipts_rows, start=2):
        _row(wr, i, values, receipts_cols)
        path = values[-1]
        if path.startswith("Чеки/"):
            cell = wr.cell(row=i, column=len(receipts_cols))
            cell.hyperlink = path
            cell.font = Font(color="0563C1", underline="single")

    xlsx = io.BytesIO()
    wb.save(xlsx)

    stem = f"{_safe(project['name'])}_{_safe(period.replace(' – ', '-'))}"
    if contractor is not None:
        stem += f"_{_safe(contractor)}"
    if not zip_files:
        return ExportResult(f"{stem}.xlsx", xlsx.getvalue(), len(shifts), grand[-1], missing)

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(f"{stem}.xlsx", xlsx.getvalue())
        for path, content in zip_files:
            zf.writestr(path, content)
    return ExportResult(f"{stem}.zip", buf.getvalue(), len(shifts), grand[-1], missing)
