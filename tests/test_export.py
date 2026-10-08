import io
import zipfile

from openpyxl import load_workbook

from bot.export import build_export


async def make_db(db):
    pid = await db.create_project(1, "Сериал «Тест»", shift_rate=5000, shift_hours=12,
                                  overtime_rate=500, km_included=100, km_rate=10)
    for uid, name in [(10, "Иванов"), (11, "Петров"), (12, "Сидоров")]:
        await db.upsert_user(uid, name, None)
    m1 = await db.add_member(pid, 10, "Иванов Иван", "Camry А001АА", "ИП Иванов")
    m2 = await db.add_member(pid, 11, "Петров Пётр", "Sprinter В002ВВ", "ООО Транс")
    m3 = await db.add_member(pid, 12, "Сидоров Сидор", "Kia С003СС", "ООО Транс")
    await db.update_member(m3, "shift_rate", 7000)
    return db, pid, (m1, m2, m3)


async def add_shift(db, member_id, day, start, end, km, status="approved", expenses=()):
    sid = await db.create_shift(member_id, f"2026-10-{day:02d}",
                                start_at=f"2026-10-{day:02d}T{start}:00+03:00",
                                end_at=f"2026-10-{day:02d}T{end}:00+03:00",
                                km=km, status=status)
    for kind, amount, file_id in expenses:
        await db.add_expense(sid, kind, amount, file_id, "photo" if file_id else None)
    return sid


async def test_export_groups_by_contractor_and_packs_receipts(db):
    db, pid, (m1, m2, m3) = await make_db(db)
    await add_shift(db, m1, 1, "08:00", "21:30", 150, expenses=[("toll", 400, "F1"), ("parking", 200, None)])
    await add_shift(db, m2, 1, "08:00", "20:00", 80, status="submitted")
    await add_shift(db, m3, 2, "07:00", "19:00", 0, expenses=[("fuel", 3000, "F2")])
    await add_shift(db, m3, 3, "07:00", "19:00", 0, status="draft")  # черновик не попадает
    await add_shift(db, m1, 20, "07:00", "19:00", 0)  # вне периода

    async def fetch(file_id, kind):
        return f"img-{file_id}".encode(), ".jpg"

    res = await build_export(db, pid, "2026-10-01", "2026-10-10", fetch_file=fetch)
    # Иванов: 5000 + 2ч*500 + 50км*10 + 600 = 7100; Петров: 5000; Сидоров: 7000 + 3000
    assert res.shifts == 3
    assert res.total == 7100 + 5000 + 10000
    assert res.missing_receipts == 1
    assert res.filename.endswith(".zip")

    zf = zipfile.ZipFile(io.BytesIO(res.data))
    names = zf.namelist()
    xlsx_name = next(n for n in names if n.endswith(".xlsx"))
    receipts = [n for n in names if n.startswith("Чеки/")]
    assert len(receipts) == 2
    wb = load_workbook(io.BytesIO(zf.read(xlsx_name)))
    assert wb.sheetnames == ["Сводка", "ИП Иванов", "ООО Транс", "Чеки"]
    summary = wb["Сводка"]
    rows = list(summary.iter_rows(min_row=7, values_only=True))
    assert rows[0][0] == "ИП Иванов" and rows[0][-1] == 7100
    assert rows[1][0] == "ООО Транс" and rows[1][1] == 2 and rows[1][-1] == 15000
    assert rows[2][0] == "ИТОГО" and rows[2][-1] == 22100
    links = [c.hyperlink.target for c in wb["Чеки"]["G"] if c.hyperlink]
    assert sorted(links) == sorted(receipts)


async def test_export_single_contractor_without_receipts_is_plain_xlsx(db):
    db, pid, (m1, m2, m3) = await make_db(db)
    await add_shift(db, m2, 1, "08:00", "20:00", 0)
    await add_shift(db, m1, 1, "08:00", "20:00", 0)
    res = await build_export(db, pid, contractor="ООО Транс")
    assert res.filename.endswith(".xlsx")
    assert res.shifts == 1 and res.total == 5000
    wb = load_workbook(io.BytesIO(res.data))
    assert wb.sheetnames == ["Сводка", "ООО Транс", "Чеки"]
